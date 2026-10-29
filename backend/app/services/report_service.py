"""
Rapport mensuel d'exploitation, en PDF et DOCX.

D'OÙ VIENNENT LES CHIFFRES. Le rapport est construit à partir des MÊMES
services que le tableau de bord — jamais de requêtes SQL parallèles écrites
pour l'occasion. Un chiffre imprimé et un chiffre affiché à l'écran qui
divergent d'un point sont le meilleur moyen de faire perdre toute confiance
dans les deux.

DEUX SOURCES. Les agrégats journaliers (`kpi_daily`) donnent la
disponibilité du parc, ses moyennes et le classement des sites. Les tickets
iTop du mois, lus à la demande comme sur l'écran SLA, donnent les volumes,
les délais (MTTO, MTTR), le respect des engagements, les causes (code de
résolution) et les dépassements. Si iTop ne répond pas, le rapport est
quand même produit et le dit dans les sections concernées.

C'est un rapport d'EXPLOITATION, pas un journal d'événements. La distinction
est assumée, et elle est écrite dans le document lui-même.
"""
from __future__ import annotations

import logging
import os
from datetime import datetime, timezone

from sqlalchemy.orm import Session

from app.core.config import REPORT_ORGANISATION, REPORT_OUTPUT_DIR
from app.services import incident_service, kpi_service, live_service

logger = logging.getLogger(__name__)

MONTHS_FR = (
    "janvier", "février", "mars", "avril", "mai", "juin",
    "juillet", "août", "septembre", "octobre", "novembre", "décembre",
)


def month_label(month: int, year: int) -> str:
    return f"{MONTHS_FR[month - 1]} {year}"


# Codes de résolution iTop (UserRequest / Incident ITIL), comme à l'écran
# (frontend/src/lib/itop.js).
RESOLUTION_CODES = {
    "assistance": "Assistance",
    "bug fixed": "Correction d'anomalie",
    "hardware repair": "Réparation matérielle",
    "other": "Autre",
    "software patch": "Correctif logiciel",
    "system update": "Mise à jour système",
    "training": "Formation",
}

SEVERITY_LABELS = {
    "critical": "Critique",
    "high": "Majeure",
    "medium": "Moyenne",
    "low": "Mineure",
    "info": "Information",
    "unknown": "Inconnue",
}


async def collect_report_data(db: Session, month: int, year: int) -> dict:
    """Rassemble les données du rapport depuis les services du dashboard."""
    monthly = kpi_service.monthly_summary(db, year, month)
    period_start, period_end = kpi_service.month_bounds(year, month)
    start_at = datetime.combine(period_start, datetime.min.time(), tzinfo=timezone.utc)
    end_at = datetime.combine(period_end, datetime.min.time(), tzinfo=timezone.utc)

    sla = incidents = None
    itop_error = None
    try:
        sla = await incident_service.sla_by_group(
            db, "region", start_at=start_at, end_at=end_at
        )
        incidents = await incident_service.incident_kpis(
            db, start_at=start_at, end_at=end_at
        )
    except (incident_service.IncidentsUnavailable, live_service.SnapshotUnavailable) as exc:
        # Le rapport reste produit : la disponibilité et les sites viennent
        # de kpi_daily. Les sections iTop disent pourquoi elles sont vides.
        logger.warning("Rapport %02d/%d : tickets iTop indisponibles (%s)", month, year, exc)
        itop_error = "Tickets iTop indisponibles au moment de la génération."

    return {
        "organisation": REPORT_ORGANISATION,
        "generated_at": datetime.now(timezone.utc),
        "period": month_label(month, year),
        "month": month,
        "year": year,
        "monthly": monthly,
        "trend": kpi_service.monthly_trend(db, year, month),
        "sites": kpi_service.sites_ranking(
            db, limit=15, start_at=period_start, end_at=period_end
        ),
        "sla": sla,
        "incidents": incidents,
        "itop_error": itop_error,
    }


def _duration(minutes) -> str:
    """Durée lisible : 47 min, 2 h 15, 3 j 4 h."""
    if minutes is None:
        return "—"
    total = max(0, round(float(minutes)))
    if total < 60:
        return f"{total} min"
    if total < 1440:
        hours, rest = divmod(total, 60)
        return f"{hours} h {rest:02d}" if rest else f"{hours} h"
    days, rest = divmod(total, 1440)
    hours = rest // 60
    return f"{days} j {hours} h" if hours else f"{days} j"


def _pct(value) -> str:
    return "—" if value is None else f"{value:.1f} %"


def incident_rows(data: dict) -> list[list[str]]:
    """Section 2 : volumes et délais des tickets iTop du mois."""
    totals = (data["sla"] or {}).get("totals") or {}
    kpis = data["incidents"] or {}
    return [
        ["Incidents ouverts dans le mois", _fmt(kpis.get("total"))],
        ["Incidents résolus", _fmt(kpis.get("resolved"))],
        ["Taux de résolution", _pct(kpis.get("resolution_rate_pct"))],
        ["Délai moyen de prise en charge (MTTO)", _duration(totals.get("mtto_minutes"))],
        ["Délai moyen de résolution (MTTR)", _duration(totals.get("mttr_minutes"))],
        ["Respect des SLA", _pct(totals.get("compliance_pct"))],
        ["Incidents hors délai", _fmt(totals.get("breached"))],
        [
            f"Équipements récurrents (≥ {kpis.get('recurrent_threshold', 3)} tickets)",
            _fmt(kpis.get("recurrent_nodes")),
        ],
        ["Incidents résolus sans cause identifiée", _pct(kpis.get("without_cause_pct"))],
        ["Incidents rouverts", _fmt(kpis.get("reopened"))],
    ]


def severity_rows(data: dict) -> list[list[str]]:
    """Section 3 : engagements par gravité."""
    sla = data["sla"] or {}
    targets = sla.get("targets") or {}
    return [
        [
            SEVERITY_LABELS.get(row["severity"], row["severity"]),
            _fmt(row.get("incidents")),
            _duration(row.get("mtto_minutes")),
            _duration(row.get("mttr_minutes")),
            _duration((targets.get(row["severity"]) or {}).get("ttr")),
            _pct(row.get("compliance_pct")),
        ]
        for row in sla.get("by_severity") or []
    ]


def cause_rows(data: dict) -> list[list[str]]:
    """Section 5 : causes retenues à la clôture des tickets."""
    return [
        [RESOLUTION_CODES.get(row["cause"], row["cause"]), str(row["count"]), _pct(row["share_pct"])]
        for row in (data["incidents"] or {}).get("causes") or []
    ]


def breach_rows(data: dict) -> list[list[str]]:
    """Section 6 : les plus longs dépassements d'objectif."""
    return [
        [
            row.get("ref") or "—",
            row.get("node_name") or "—",
            SEVERITY_LABELS.get(row["severity"], row["severity"]),
            _duration(row["ttr_minutes"]) + (" (en cours)" if row.get("open") else ""),
            _duration(row["target_minutes"]),
            _duration(row["overrun_minutes"]),
        ]
        for row in (data["incidents"] or {}).get("breaches") or []
    ]


def _no_cause_note(data: dict) -> str:
    if data["itop_error"]:
        return data["itop_error"]
    if not (data["incidents"] or {}).get("cause_available"):
        return "L'instance iTop ne publie pas de code de résolution : causes non disponibles."
    return "Aucune cause renseignée dans iTop sur les tickets résolus du mois."


METHOD_NOTE = (
    "Méthode : disponibilité et sites viennent des relevés de collecte du NOC "
    "(moyennes journalières, maintenances planifiées exclues) ; incidents, "
    "délais, causes et dépassements viennent des tickets iTop ouverts dans le "
    "mois. Le détail événement par événement reste consultable dans les outils "
    "de supervision et dans iTop, qui en sont les sources de référence."
)


def _ensure_output_dir() -> str:
    os.makedirs(REPORT_OUTPUT_DIR, exist_ok=True)
    return REPORT_OUTPUT_DIR


def _filename(month: int, year: int, extension: str) -> str:
    return os.path.join(
        _ensure_output_dir(), f"rapport-noc-{year}-{month:02d}.{extension}"
    )


def _fmt(value, suffix: str = "", default: str = "—") -> str:
    """Valeur formatée, avec un tiret cadratin quand elle est absente.

    Un « 0 » à la place d'une donnée manquante est un mensonge : il se lit
    « aucun incident » là où la vérité est « nous ne savons pas ». Le tiret
    dit qu'on ne sait pas.
    """
    if value is None:
        return default
    if isinstance(value, float):
        return f"{value:.2f}{suffix}"
    return f"{value}{suffix}"


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def render_pdf(data: dict) -> str:
    from fpdf import FPDF

    month, year = data["month"], data["year"]

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Les polices cœur de fpdf2 sont en latin-1 : le texte est transcodé
    # explicitement, sinon chaque accent français lève une UnicodeEncodeError
    # au moment de l'écriture.
    def _txt(value) -> str:
        return str(value).encode("latin-1", "replace").decode("latin-1")

    # fpdf2 >= 2.7 déprécie l'argument `ln=` de cell() : selon la version il
    # est ignoré, le curseur reste alors collé à la marge droite et l'appel
    # multi_cell() suivant échoue avec « Not enough horizontal space ». On
    # pilote donc le retour à la ligne à la main.
    def _newline(height: float = 6) -> None:
        pdf.ln(height)
        pdf.set_x(pdf.l_margin)

    def _line(text: str, height: float = 6) -> None:
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(0, height, _txt(text))

    def _row(cells, height: float = 6, border: int = 0) -> None:
        pdf.set_x(pdf.l_margin)
        for width, value, align in cells:
            pdf.cell(width, height, _txt(value), border=border, align=align)
        _newline(height)

    def _section(title: str) -> None:
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 12)
        _line(title, 8)
        pdf.set_font("Helvetica", "", 10)

    pdf.set_font("Helvetica", "B", 16)
    _line(f"Rapport d'exploitation NOC — {data['period']}", 10)
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(110, 110, 110)
    _line(
        f"{data['organisation']} — généré le "
        f"{data['generated_at'].strftime('%d/%m/%Y à %H:%M UTC')}"
    )
    pdf.set_text_color(0, 0, 0)

    # -- 1. Indicateurs ---------------------------------------------------
    _section("1. Indicateurs du mois")
    current = data["monthly"]["current"]
    previous = data["monthly"]["previous"]
    delta = data["monthly"]["delta"]

    _row([(90, "Indicateur", "L"), (35, "Mois", "R"), (35, "Précédent", "R"), (30, "Écart", "R")], border=1)
    for key, label, suffix in (
        ("availability_pct", "Disponibilité du parc", " %"),
        ("avg_alerts", "Alertes actives (moyenne)", ""),
        ("avg_critical", "Alertes critiques (moyenne)", ""),
        ("avg_down", "Équipements en panne (moyenne)", ""),
    ):
        _row(
            [
                (90, label, "L"),
                (35, _fmt(current.get(key), suffix), "R"),
                (35, _fmt(previous.get(key), suffix), "R"),
                (30, _fmt(delta.get(key)), "R"),
            ]
        )

    if not data["monthly"]["reliable"]:
        pdf.set_font("Helvetica", "I", 9)
        _line(
            f"Attention : seulement {current['days_with_data']} jour(s) de "
            "données sur ce mois. Les moyennes ci-dessus ne sont pas "
            "comparables à celles d'un mois complet.",
            5,
        )
        pdf.set_font("Helvetica", "", 10)

    # -- 2. Incidents -----------------------------------------------------
    _section("2. Incidents iTop du mois")
    if data["itop_error"]:
        _line(data["itop_error"])
    else:
        for label, value in incident_rows(data):
            _row([(120, label, "L"), (60, value, "L")])

    # -- 3. Respect des engagements ---------------------------------------
    _section("3. Respect des engagements de service")
    rows = severity_rows(data)
    if rows:
        _row([(30, "Gravité", "L"), (25, "Incidents", "R"), (30, "MTTO", "R"),
              (30, "MTTR", "R"), (35, "Objectif MTTR", "R"), (30, "Respect", "R")], border=1)
        for row in rows:
            _row(list(zip((30, 25, 30, 30, 35, 30), row, ("L", "R", "R", "R", "R", "R"))))
    else:
        _line(data["itop_error"] or "Aucun incident iTop sur la période.")

    # -- 4. Sites ---------------------------------------------------------
    _section("4. Sites les plus affectés")
    if data["sites"]:
        _row([(70, "Site", "L"), (40, "Disponibilité", "R"),
              (40, "Alertes (moy.)", "R"), (40, "En panne (moy.)", "R")], border=1)
        for site in data["sites"]:
            _row(
                [
                    (70, site["site"] or "—", "L"),
                    (40, _fmt(site.get("availability_pct"), " %"), "R"),
                    (40, _fmt(site.get("avg_alerts")), "R"),
                    (40, _fmt(site.get("avg_down")), "R"),
                ]
            )
    else:
        _line("Aucune donnée par site sur la période.")

    # -- 5. Causes --------------------------------------------------------
    _section("5. Causes (code de résolution iTop)")
    rows = cause_rows(data)
    if rows:
        for cause, count, share in rows:
            _row([(120, cause, "L"), (30, count, "R"), (30, share, "R")])
    else:
        _line(_no_cause_note(data))

    # -- 6. Dépassements --------------------------------------------------
    _section("6. Dépassements d'objectif")
    rows = breach_rows(data)
    if rows:
        _row([(25, "Ticket", "L"), (55, "Équipement", "L"), (22, "Gravité", "L"),
              (30, "Résolution", "R"), (28, "Objectif", "R"), (30, "Dépassement", "R")], border=1)
        for row in rows:
            _row(list(zip((25, 55, 22, 30, 28, 30), row, ("L", "L", "L", "R", "R", "R"))), 5)
    else:
        _line(data["itop_error"] or "Aucun dépassement sur la période.")

    # -- Note de méthode --------------------------------------------------
    pdf.ln(4)
    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(110, 110, 110)
    _line(METHOD_NOTE, 4)
    pdf.set_text_color(0, 0, 0)

    path = _filename(month, year, "pdf")
    pdf.output(path)
    logger.info("Rapport PDF généré : %s", path)
    return path


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def render_docx(data: dict) -> str:
    from docx import Document
    from docx.shared import Pt

    month, year = data["month"], data["year"]
    document = Document()

    document.add_heading(f"Rapport d'exploitation NOC — {data['period']}", level=0)
    subtitle = document.add_paragraph(
        f"{data['organisation']} — généré le "
        f"{data['generated_at'].strftime('%d/%m/%Y à %H:%M UTC')}"
    )
    subtitle.runs[0].font.size = Pt(9)

    def _table(headers: list[str], rows: list[list[str]]) -> None:
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header
        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = str(value)

    document.add_heading("1. Indicateurs du mois", level=1)
    current, previous, delta = (
        data["monthly"]["current"],
        data["monthly"]["previous"],
        data["monthly"]["delta"],
    )
    _table(
        ["Indicateur", "Mois", "Précédent", "Écart"],
        [
            [label, _fmt(current.get(key), suffix), _fmt(previous.get(key), suffix), _fmt(delta.get(key))]
            for key, label, suffix in (
                ("availability_pct", "Disponibilité du parc", " %"),
                ("avg_alerts", "Alertes actives (moyenne)", ""),
                ("avg_critical", "Alertes critiques (moyenne)", ""),
                ("avg_down", "Équipements en panne (moyenne)", ""),
            )
        ],
    )
    if not data["monthly"]["reliable"]:
        warning = document.add_paragraph(
            f"Attention : seulement {current['days_with_data']} jour(s) de données "
            "sur ce mois. Les moyennes ne sont pas comparables à un mois complet."
        )
        warning.runs[0].italic = True

    document.add_heading("2. Incidents iTop du mois", level=1)
    if data["itop_error"]:
        document.add_paragraph(data["itop_error"])
    else:
        _table(["Indicateur", "Valeur"], incident_rows(data))

    document.add_heading("3. Respect des engagements de service", level=1)
    rows = severity_rows(data)
    if rows:
        _table(["Gravité", "Incidents", "MTTO", "MTTR", "Objectif MTTR", "Respect"], rows)
    else:
        document.add_paragraph(data["itop_error"] or "Aucun incident iTop sur la période.")

    document.add_heading("4. Sites les plus affectés", level=1)
    if data["sites"]:
        _table(
            ["Site", "Disponibilité", "Alertes (moy.)", "En panne (moy.)"],
            [
                [
                    site["site"] or "—",
                    _fmt(site.get("availability_pct"), " %"),
                    _fmt(site.get("avg_alerts")),
                    _fmt(site.get("avg_down")),
                ]
                for site in data["sites"]
            ],
        )
    else:
        document.add_paragraph("Aucune donnée par site sur la période.")

    document.add_heading("5. Causes (code de résolution iTop)", level=1)
    rows = cause_rows(data)
    if rows:
        _table(["Cause", "Occurrences", "Part"], rows)
    else:
        document.add_paragraph(_no_cause_note(data))

    document.add_heading("6. Dépassements d'objectif", level=1)
    rows = breach_rows(data)
    if rows:
        _table(["Ticket", "Équipement", "Gravité", "Résolution", "Objectif", "Dépassement"], rows)
    else:
        document.add_paragraph(data["itop_error"] or "Aucun dépassement sur la période.")

    note = document.add_paragraph(METHOD_NOTE)
    note.runs[0].font.size = Pt(8)
    note.runs[0].italic = True

    path = _filename(month, year, "docx")
    document.save(path)
    logger.info("Rapport DOCX généré : %s", path)
    return path


async def generate(db: Session, month: int, year: int, output_format: str) -> str:
    if output_format not in ("pdf", "docx"):
        raise ValueError(f"Format inconnu : {output_format}")
    data = await collect_report_data(db, month, year)
    return render_pdf(data) if output_format == "pdf" else render_docx(data)
